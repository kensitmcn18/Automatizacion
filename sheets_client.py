import requests
from deep_translator import GoogleTranslator
import pandas as pd
from collections import defaultdict
from datetime import datetime
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
import time
from typing import Dict, Set, List, Optional


CLIENTOKEN = "E0D439EE522F44368DC78E1BFB03710C-D24FB11DBE31D4621C4817E028D9E1D"
ACCESSTOKEN = "C66EF7B239D24632943D115EDE9CB810-EA00F8FD8294692C940F6B5A8F9453D"
HEADERS = {"Content-Type": "application/json"}
START_DATE = "2025-05-06T01:00:00Z"
END_DATE = "2025-05-08T23:00:00Z"

def adjust_column_widths(worksheet, df):
    column_config = {
    
                'columnas_cortas': {
                    'columns': [
                        'Numero', 'Estado', 'Recuento noches', 'Cantidad de Personas', 
                        'Recuento camas', 'Numero habitación', 'Precio', 'Importe total',
                        'Tipo tarjeta', 'Nacionalidad'
                    ],
                        'factor': 1.3,
                        'margin': 3,
                        'min_width': 12,
                        'max_width': 20},

                'columnas_largas': {
                    'columns': [
                        'Correo', 'Direccion', 'Productos', 'RecibeCorreosMarketing',
                        'Tipo habitación'
                    ],
                        'factor': 1.2,
                        'margin': 5,
                        'min_width': 25,
                        'max_width': 60},
    
                'columnas_fechas': {
                    'columns': [
                        'Fecha de creación', 'Fecha de actualización', 'Fecha llegada', 
                        'Fecha salida', 'Fecha vencimiento'
                    ],
                        'factor': 1.1,
                        'margin': 2,
                        'min_width': 18,
                        'max_width': 25},
    
    
                'columnas_texto_medio': {
                    'columns': [
                        'Grupo', 'Apellido', 'Nombre', 'Teléfono', 'tarifa', 'Numero de tarjeta',
                        'motivo de la reserva'
                    ],
                        'factor': 1.4,
                        'margin': 4,
                        'min_width': 15,
                        'max_width': 35}
                    
                            }
    default_config = {
                        'factor': 1.5,
                        'margin': 5,
                        'min_width': 15,
                        'max_width': 40
                    }
    
    def get_column_config(column_name, column_config, default_config):
        for category, config in column_config.items():
            if column_name in config['columns']:
                return config
        return default_config 

    def calculate_column_width(max_length, config):
        calculated_width = max_length * config['factor']+ config['margin']
        return max(min(calculated_width, config['max_width']), config['min_width'])
    for i, column_name in enumerate(df.columns):
        max_length = len(column_name)
        column_letter = get_column_letter(i + 1)
        for row_data in df[column_name]:
            cell_value = str(row_data) if row_data is not None else ""
            max_length = max(max_length, len(cell_value))

        config = get_column_config(column_name, column_config, default_config)
        adjusted_width = calculate_column_width(max_length, config)
        worksheet.column_dimensions[column_letter].width = adjusted_width

def get_mews_reservations():
    url_reservations = "https://api.mews-demo.com/api/connector/v1/reservations/getAll/2023-06-06"
    
    
    payload_reservations = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "States": [
            "Inquired",
            "Confirmed",
            "Optional",
            "Started",
            "Processed",
            "Requested"
        ],
        "UpdatedUtc": {"StartUtc": START_DATE, 
                        "EndUtc": END_DATE},
        "Limitation": {"Count": 30}
    }
    
    response = requests.post(url_reservations, json=payload_reservations, headers=HEADERS)
    
    if response.status_code == 200:
        data = response.json()
        reservations = data.get("Reservations", [])

        if reservations:
            
            reservation_ids = [reserva["Id"] for reserva in reservations]
            service_ids = {reserva.get("ServiceId") for reserva in reservations if reserva.get("ServiceId")}
            group_ids = {reserva.get("GroupId") for reserva in reservations if reserva.get("GroupId")}
            account_ids = {reserva.get("AccountId") for reserva in reservations if reserva.get("AccountId")}
            counts = [person["Count"] for reservation in data.get("Reservations", []) for person in reservation.get("PersonCounts", [])]
            productos_por_reserva = get_products_by_reservation(reservation_ids, service_ids)
            rate_ids = {res.get("RateId") for res in reservations if res.get("RateId")}
            rate_prices = get_rates_pricing_batch(rate_ids)
            
                            
            group_names = get_group_names(group_ids)
            customer_info = get_customer_names(account_ids)
            rate_name = get_rates_names()
            
        
            reservas_filtradas = []    
            for reserva in reservations:
                service_id = reserva.get("ServiceId")
                reservation_ids = [reserva["Id"]for reserva in reservations]
                reservation_id = reserva["Id"]
                
                               
                if not reserva.get("CreditCardId"):
                    tipo_tarjeta, numero_tarjeta, expiracion_tarjeta = "", "", ""
                else:
                    tipo_tarjeta, numero_tarjeta, expiracion_tarjeta = get_credit_card()
                    
           
            reservas_filtradas = [
                {
                    "Numero": reserva.get("Number", ""),
                    "Grupo": group_names.get(reserva.get("GroupId"), ""),
                    "Apellido": customer_info.get(reserva.get("AccountId"), {}).get("Apellido", ""),
                    "Nombre": customer_info.get(reserva.get("AccountId"), {}).get("Nombre", ""),
                    "Correo": customer_info.get(reserva.get("AccountId"), {}).get("Correo", ""),
                    "Teléfono": customer_info.get(reserva.get("AccountId"), {}).get("Telefono", ""),
                    "Direccion": customer_info.get(reserva.get("AccountId"), {}).get("Direccion", ""),
                    "Nacionalidad": customer_info.get(reserva.get("AccountId"), {}).get("Nacionalidad", ""),
                    "RecibeCorreosMarketing": customer_info.get(reserva.get("AccountId"), {}).get("Enviar correos electronicos de Marketing", ""),
                    "Estado": reserva.get("State", ""), 
                    "Fecha de creación": reserva.get("CreatedUtc", ""),
                    "Fecha de actualización": reserva.get("UpdatedUtc", ""),
                    "Fecha llegada": reserva.get("ScheduledStartUtc", ""),
                    "Fecha salida": reserva.get("ScheduledEndUtc", ""),
                    "Recuento noches": calcular_noches(reserva),                    
                    "Cantidad de Personas": sum(person["Count"] for person in reserva.get("PersonCounts", []) if "Count" in person),
                    "Recuento camas": calcular_camas(reserva),
                    "Tipo habitación": get_room_name(reserva.get("ServiceId"), reserva.get("RequestedResourceCategoryId")),
                    "Numero habitación": get_room_number(reserva.get("AssignedResourceId")),
                    "tarifa": rate_name.get(reserva.get("RateId")),
                    "Productos": productos_por_reserva.get(reserva.get("Id"),""),
                    "Precio": rate_prices.get(reserva.get("RateId"), ""),
                    "Importe total": calcular_importe_total(rate_prices.get(reserva.get("RateId")), calcular_noches(reserva)), 
                    "Tipo tarjeta": tipo_tarjeta,
                    "Numero de tarjeta": numero_tarjeta,
                    "Fecha vencimiento": expiracion_tarjeta,
                    "motivo de la reserva": reserva.get("Purpose", "")
                }
                for reserva in reservations

            ]
              
        df = pd.DataFrame(reservas_filtradas)
        output_filename = "reservas_mews.xlsx"
        df.to_excel(output_filename, index=False, sheet_name="Reservas")
        workbook = load_workbook(output_filename)
        worksheet = workbook["Reservas"]
        adjust_column_widths(worksheet, df)

        workbook.save(output_filename)
        print(f"Archivo '{output_filename}' guardado exitosamente")
        categories_used = {}
            
        
def get_group_names(group_ids):
    if not group_ids:
        return {}
    
    url_groups = "https://api.mews-demo.com/api/connector/v1/reservationGroups/getAll"
    
    
    payload_groups = {
       "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "UpdatedUtc": {"StartUtc": START_DATE, 
                        "EndUtc": END_DATE},
        "GroupIds": list(group_ids),
        "Limitation": {"Count": 30}
    }
    
    response = requests.post(url_groups, json=payload_groups, headers=HEADERS)
    
    if response.status_code == 200:
        data = response.json()
        return {group["Id"]: group["Name"] for group in data.get("ReservationGroups", [])}
    else:
        print(f"Error {response.status_code} al obtener nombres de grupos: {response.text}")
        return {}
def translate_country(country_name):
    
    if not country_name:
        return ""
    try:
        return GoogleTranslator(source="en", target="es").translate(country_name)
    except Exception:
        return country_name 

def get_countries():
    url_countries = "https://api.mews-demo.com/api/connector/v1/countries/getAll"
    
    payload_countries = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN}
    response = requests.post(url_countries, json=payload_countries, headers=HEADERS)
    
    if response.status_code == 200:
        countries_data = response.json().get("Countries", [])
        country_dict = {country["Code"]: translate_country(country["EnglishName"]) for country in countries_data}
        
        return country_dict
    return {}

def get_customer_names(account_ids):
    if not account_ids:
        return {}
    
    url_customers = "https://api.mews-demo.com/api/connector/v1/customers/getAll"

    payload_customers = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "CustomerIds": list(account_ids),
        "Extent": {
                "Customers": True,
                "Addresses": True},
        "Limitation": {"Count": 30}, 
        
    }

    response = requests.post(url_customers, json=payload_customers, headers=HEADERS)

    if response.status_code == 200:
        customers_data = response.json().get("Customers", [])
        countries_data = get_countries()

        customer_info = {}
        for customer in customers_data:
            customer_id = customer.get("Id", "")
            nationality_code = customer.get("NationalityCode", "")
            country_name = countries_data.get(nationality_code)
    
            
            address = customer.get("Address") or {}
            address_parts = [
                address.get("Line1"), 
                address.get("Line2"), 
                address.get("City"), 
                address.get("PostalCode"), 
                countries_data.get(address.get("CountryCode"), address.get("CountryCode"))
            ]
            formatted_address = ", ".join(filter(None, address_parts))
            send_marketing_emails = "No"
            if "Options" in customer and "SendMarketingEmails" in customer["Options"]:
                send_marketing_emails = "Yes"
           
            customer_info[customer_id] = {
                "Nombre": customer.get("FirstName", ""),
                "Apellido": customer.get("LastName", ""),
                "Correo": customer.get("Email",""),
                "Telefono": customer.get("Phone",""),
                "Direccion": formatted_address,
                "Nacionalidad": country_name,
                "Enviar correos electronicos de Marketing": send_marketing_emails
                
            }
        return customer_info
    else:
        print(f"Error {response.status_code} al obtener clientes: {response.text}")
        return {}

def get_rates_names():
    url_rates_names = "https://api.mews-demo.com/api/connector/v1/rates/getAll"

    payload_rates_names = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN
    }
    response = requests.post(url_rates_names, json=payload_rates_names, headers=HEADERS)

    if response.status_code == 200:
        rates_data = response.json().get("Rates", [])
        return {rate["Id"]: rate["Name"] for rate in rates_data if "Id" in rate and "Name" in rate}
    else:
        print(f"Error al obtener nombres de tarifas: {response.status_code} - {response.text}")
        return {}

def get_rates_pricing_batch(rate_ids: Set[str], max_retries: int = 5):
    if not rate_ids:
        return {}
    
    url_rates = "https://api.mews-demo.com/api/connector/v1/rates/getPricing"
    rate_prices = {}
   
    for rate_id in rate_ids: 
        if not rate_id:
            continue
    payload_rates = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN, 
        "RateId": rate_id,
        "StartUtc": START_DATE,
        "EndUtc": END_DATE    
    }
    wait_time = 1
    success = False
    for attempt in range(max_retries):
        try:
            response = requests.post(url_rates, json=payload_rates, headers=HEADERS)
    
            if response.status_code == 200:
                data = response.json()
                currency_map = {"GBP": "£", "USD": "$", "EUR": "€"}

                for item in data.get("BaseAmountPrices", []):
                    currency = item.get('Currency', '')
                    gross_value = item.get('GrossValue', 0)
                    symbol = currency_map.get(currency, currency)
                    rate_prices[item.get("RateId")] = f"{symbol} {gross_value}"
                success = True
                break
            elif response.status_code == 429:
                retry_after = int(response.headers.get("Retry-After", wait_time))
                print(f"⏳ Rate limit alcanzado para RateId {rate_id}. Esperando {retry_after} segundos...")
                time.sleep(retry_after)
                wait_time = min(wait_time * 2, 60)
            else: 
                print(f"❌ Error {response.status_code} para RateId {rate_id}: {response.text}")
                break
        except requests.exceptions.RequestException as e: 
            print(f"Error de conexion para RateId {rate_id}:{e}")
            break
    if not success:
        print(f" Fallo obtener precio para RateId{rate_id}despues de {max_retries} intentos")

    time.sleep (0.5)
    return rate_prices


def get_credit_card(Credit_card): 
    url_card = "https://api.mews-demo.com/api/connector/v1/creditCards/getAll"
    payload_card = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "UpdatedUtc": {"StartUtc": START_DATE, 
                        "EndUtc": END_DATE}
    }
    
    response = requests.post(url_card, json=payload_card, headers=HEADERS)
    if response.status_code == 200: 
        card_data = response.json()
        if not card_data or "CreditCards" not in card_data or not card_data["CreditCards"]:
            return "", "", ""  

        for card in card_data["CreditCards"]:
            if card.get("ObfuscatedNumber"):  
                return card.get("Type", ""), card.get("ObfuscatedNumber", ""), card.get("Expiration", "")

    return "", "", ""

          
def get_room_name(service_id, resource_category_id):
    url = "https://api.mews-demo.com/api/connector/v1/resourceCategories/getAll"
    
    
    payload = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN, 
        "ResourceCategoryIds": [resource_category_id],
        "ServiceIds": [service_id]
    }
    
    response = requests.post(url, json=payload, headers=HEADERS)
    
    if response.status_code == 200:
        data = response.json()
        if data and "ResourceCategories" in data and len(data["ResourceCategories"]) > 0:
            names = data["ResourceCategories"][0].get("Names", {})
            return names.get("en-US", "Tipo no encontrado")
        return "Datos no encontrados"
    else:
        return "Información no encontrada"

def get_room_number(assigned_resource_id):
    url_resources = "https://api.mews-demo.com/api/connector/v1/resources/getAll"
  
    
    payload = {
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "ResourceIds": [assigned_resource_id]
    }
    
    response = requests.post(url_resources, json=payload, headers=HEADERS)
    
    if response.status_code == 200:
        data = response.json()
        if data and "Resources" in data and len(data["Resources"]) > 0:
            return data["Resources"][0].get("Name", "Número no encontrado")
        return "Datos no encontrados"
    else:
        return "Información no encontrada"
    
def get_products_by_reservation(reservation_ids, service_ids):
    url_items = "https://api.mews-demo.com/api/connector/v1/orderItems/getAll"
    payload ={
        "ClientToken": CLIENTOKEN,
        "AccessToken": ACCESSTOKEN,
        "ServiceOrderIds": list(reservation_ids),
        "ServiceIds": list(service_ids),
        "Limitation": {"Count": 40}   
    }
    
    response = requests.post(url_items, json=payload, headers=HEADERS)
    
    
    if response.status_code == 200:
        order_items =  response.json().get("OrderItems", [])
        productos_por_reserva = defaultdict(list)
        valor = ""
        for item in order_items:
            res_id = item.get("ServiceOrderId")
            productType = item.get("Type", "").strip()
            if productType == "ProductOrder":
                nombre_producto = item.get("BillingName", "").strip()
                cantidad = item.get("UnitCount",1)
                valor = str(cantidad) + " x " + nombre_producto
            if res_id and valor != "":
                productos_por_reserva[res_id].append(valor)

        return {
            res_id: ", ".join(sorted(set(productos)))
            for res_id, productos in productos_por_reserva.items()
        }
    else:
        print(f"Error al obtener OrderItems: {response.status_code} - {response.text}")
        return {}
    
def calcular_noches(reserva):

    try: 
        fecha_llegada_str = reserva.get("ScheduledStartUtc")
        fecha_salida_str = reserva.get("ScheduledEndUtc")
        if not fecha_llegada_str or not fecha_salida_str:
            return ""
        formatos = [
            "%Y-%m-%dT%H:%M:%SZ",      
            "%Y-%m-%dT%H:%M:%S.%fZ",   
            "%Y-%m-%dT%H:%M:%S",       
        ]
        fecha_llegada = None
        fecha_salida = None
        for formato in formatos:
            try:
                fecha_llegada = datetime.strptime(fecha_llegada_str.strip(), formato).date()
                break
            except ValueError:
                continue
        for formato in formatos:
            try:
                fecha_salida = datetime.strptime(fecha_salida_str.strip(), formato).date()
                break
            except ValueError:
                continue
        
        if fecha_llegada and fecha_salida:
            noches = (fecha_salida - fecha_llegada).days
            return noches if noches >= 0 else ""
        else:
            return ""
            
    except Exception as e:
        print(f"Error al calcular noches para reserva {reserva.get('Number')}: {e}")
        return ""
def calcular_camas(reserva):
    
    try:
        noches = calcular_noches(reserva)
        personas = sum(person["Count"] for person in reserva.get("PersonCounts", []) if "Count" in person)
        
        if noches == "" or personas == 0:
            return ""
        
        return noches * personas
    except Exception as e:
        print(f"Error al calcular persona-noches para reserva {reserva.get('Number')}: {e}")
        return ""

def calcular_importe_total(precio_str, noches):
    try:
        if not precio_str or not isinstance(noches, int):
            return ""
        simbolo, valor_str = precio_str.split()
        valor = float(valor_str.replace(',', ''))
        return f"{simbolo} {round(valor * noches, 2)}"
    except Exception as e:
        print(f"Error al calcular importe total: {e}")
        return ""



if __name__ == "__main__":
    get_mews_reservations()
